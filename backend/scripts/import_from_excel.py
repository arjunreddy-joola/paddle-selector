#!/usr/bin/env python3
"""
Import paddle recommendation data from the JOOLA Excel file.
Usage: python scripts/import_from_excel.py --file "path/to/JOOLA Paddle Selector.xlsx"

Requires: pip install openpyxl
"""
import argparse
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    sys.exit(1)

COLUMN_MAP = {
    "Q1: How often do you play?": "playFrequency",
    "Q2: What's your budget?": "budget",
    "Q3: What matters most to you?": "playPriority",
    "Q4: Do you use a two-handed backhand?": "twoHandedBackhand",
    "Q5: Are you open to upgrading if it improves performance?": "upgradeOpenness",
    "Q6: How does your current paddle grip feel?": "gripFeel",
}

VALUE_NORMALIZER = {
    "Occasionally": "occasionally",
    "Weekly": "weekly",
    "Multiple times per week": "multiple_times_per_week",
    "Under $100": "under_100",
    "$100–$200": "100_200",
    "$200+ (premium performance)": "200_plus",
    "No budget": "no_budget",
    "Power": "power",
    "Control": "control",
    "Balanced": "balance",
    "Spin": "spin",
    "Yes": "yes",
    "No": "no",
    "Sometimes": "sometimes",
    "Maybe": "maybe",
    "Small": "small",
    "Average / Just right": "average",
    "Large": "large",
    "Not sure / This is my 1st paddle": "not_sure_first_paddle",
}


def import_excel(excel_path: str) -> None:
    from src.utils.parse_recommendation_notes import parse_recommendation_notes

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb["All Combinations"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

    existing_path = Path("src/data/recommendation_rules.json")
    existing = json.loads(existing_path.read_text(encoding="utf-8"))
    rules_by_key: dict[str, dict] = {}
    for rule in existing["rules"]:
        key = "|".join(f"{f}={rule['answers'][f]}" for f in existing["lookupFields"])
        rules_by_key[key] = rule

    updated = 0
    skipped = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_data = dict(zip(headers, row))
        # Skip empty rows
        if not any(row_data.values()):
            continue

        answers: dict[str, str] = {}
        excel_answers: dict[str, str] = {}
        for col_name, field_id in COLUMN_MAP.items():
            raw_val = str(row_data.get(col_name) or "").strip()
            answers[field_id] = VALUE_NORMALIZER.get(raw_val, raw_val.lower().replace(" ", "_"))
            excel_answers[field_id] = raw_val

        key = "|".join(f"{f}={answers[f]}" for f in existing["lookupFields"])
        if key not in rules_by_key:
            print(f"  No matching rule for key: {key}")
            skipped += 1
            continue

        notes_raw = str(row_data.get("Notes") or "")
        product_list_raw = str(row_data.get("Product List (Paddle Recommendation)") or "")
        best_seller_raw = str(row_data.get("Best Seller") or "").strip()
        recommended_for_you_raw = str(row_data.get("Recommended for you") or "").strip()

        rule = rules_by_key[key]
        rule["rawNotes"] = notes_raw
        rule["rawProductList"] = product_list_raw
        rule["excelAnswers"] = excel_answers

        products = [p.strip() for p in product_list_raw.split("|") if p.strip()]
        rule["otherPaddles"] = products

        # Use the dedicated columns first, fall back to Notes parsing
        if recommended_for_you_raw:
            rule["recommendedForYou"] = {"productName": recommended_for_you_raw, "source": "notes"}
            rule["requiresMerchandisingReview"] = False
            rule["fallbackUsed"] = False
        else:
            parsed = parse_recommendation_notes(notes_raw)
            if parsed["recommendedForYou"]:
                rule["recommendedForYou"] = {"productName": parsed["recommendedForYou"], "source": "notes"}
                rule["requiresMerchandisingReview"] = False
                rule["fallbackUsed"] = False

        if best_seller_raw:
            rule["bestSeller"] = {"productName": best_seller_raw, "source": "notes"}
        else:
            parsed = parse_recommendation_notes(notes_raw)
            if parsed["bestSeller"]:
                rule["bestSeller"] = {"productName": parsed["bestSeller"], "source": "notes"}

        updated += 1

    existing_path.write_text(json.dumps(existing, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Updated {updated} rules from Excel -> {existing_path}")
    if skipped:
        print(f"Skipped {skipped} unmatched rows")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True, help="Path to Excel file")
    args = parser.parse_args()
    import_excel(args.file)
